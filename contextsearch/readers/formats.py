"""Built-in readers: PDF, DOCX, PPTX, plain text, HTML and images.

Each reader does one format and nothing else. Errors propagate and are turned
into :class:`ReadError` by :func:`readers.base.read_document`.
"""

from __future__ import annotations

import csv
import io
import json
from html.parser import HTMLParser
from pathlib import Path
from typing import TYPE_CHECKING

from ..errors import ReadError
from ..observability import get_logger
from . import ocr
from .base import Page, ParsedDocument, Reader, register

if TYPE_CHECKING:
    from ..config import Settings

log = get_logger("contextsearch.readers")

TEXT_ENCODINGS = ("utf-8-sig", "utf-16", "cp1252", "latin-1")


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in TEXT_ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# ------------------------------------------------------------------ PDF


class PdfReader(Reader):
    """PDF text through PDFium. Pages without a text layer go to OCR when enabled."""

    name = "pdf"
    extensions = ("pdf",)

    def read(self, path: Path, settings: Settings) -> ParsedDocument:
        try:
            import pypdfium2 as pdfium
        except ImportError as exc:  # pragma: no cover
            raise ReadError("pypdfium2 is not installed") from exc

        try:
            doc = pdfium.PdfDocument(str(path))
        except pdfium.PdfiumError as exc:
            raise ReadError(f"PDFium cannot open {path.name}: {exc}", {"path": str(path)}) from exc

        pages: list[Page] = []
        title = ""
        try:
            try:
                title = (doc.get_metadata_value("Title") or "").strip()
            except Exception:
                title = ""
            for index in range(len(doc)):
                page = doc[index]
                try:
                    textpage = page.get_textpage()
                    try:
                        text = textpage.get_text_range() or ""
                    finally:
                        textpage.close()
                    used_ocr = False
                    if settings.ocr and len(text.strip()) < settings.ocr_min_page_chars:
                        if ocr.available():
                            image = page.render(scale=settings.ocr_render_scale).to_pil()
                            text, _conf = ocr.image_to_text(image)
                            used_ocr = True
                        else:
                            log.warning(
                                "page has no text layer and OCR is unavailable",
                                extra_fields={"path": str(path), "page": index + 1},
                            )
                    pages.append(Page(number=index + 1, text=text, ocr=used_ocr))
                finally:
                    page.close()
        finally:
            doc.close()
        return ParsedDocument(pages=pages, title=title, metadata={"pages": len(pages)})


# ------------------------------------------------------------------ DOCX


class DocxReader(Reader):
    """Word documents: paragraphs and tables in document order, as one page."""

    name = "docx"
    extensions = ("docx",)

    def read(self, path: Path, settings: Settings) -> ParsedDocument:
        from docx import Document
        from docx.oxml.ns import qn

        doc = Document(str(path))
        parts: list[str] = []
        # Walk the body so tables land where they appear, not after all paragraphs.
        body = doc.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                text = "".join(node.text or "" for node in child.iter(qn("w:t")))
                parts.append(text)
            elif child.tag == qn("w:tbl"):
                for row in child.iter(qn("w:tr")):
                    cells = []
                    for cell in row.iter(qn("w:tc")):
                        cells.append(" ".join(node.text or "" for node in cell.iter(qn("w:t"))).strip())
                    parts.append(" | ".join(c for c in cells if c))
                parts.append("")
        title = ""
        try:
            title = (doc.core_properties.title or "").strip()
        except Exception:
            pass
        if not title:
            for p in doc.paragraphs:
                if (
                    p.style is not None
                    and p.style.name
                    and p.style.name.lower().startswith(("heading", "title"))
                    and p.text.strip()
                ):
                    title = p.text.strip()
                    break
        return ParsedDocument(pages=[Page(number=1, text="\n".join(parts))], title=title)


# ------------------------------------------------------------------ PPTX


class PptxReader(Reader):
    """PowerPoint decks: one page per slide, including notes."""

    name = "pptx"
    extensions = ("pptx",)

    def read(self, path: Path, settings: Settings) -> ParsedDocument:
        from pptx import Presentation

        pres = Presentation(str(path))
        pages: list[Page] = []
        title = ""
        for number, slide in enumerate(pres.slides, start=1):
            lines: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
                    for paragraph in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs).strip()
                        if text:
                            lines.append(text)
                if getattr(shape, "has_table", False) and shape.table is not None:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            lines.append(" | ".join(cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append("")
                    lines.append("Notes: " + notes)
            if number == 1 and lines and not title:
                title = lines[0][:120]
            pages.append(Page(number=number, text="\n".join(lines)))
        return ParsedDocument(pages=pages, title=title, metadata={"slides": len(pages)})


# ------------------------------------------------------------------ plain text


class TextReader(Reader):
    """Plain text and text-like formats. CSV rows are joined so cell values stay searchable."""

    name = "text"
    extensions = ("txt", "md", "markdown", "rst", "log", "csv", "json", "yaml", "yml", "xml", "ini", "cfg", "toml")

    def read(self, path: Path, settings: Settings) -> ParsedDocument:
        text = _read_text_file(path)
        ext = path.suffix.lower().lstrip(".")
        if ext == "csv":
            rows = list(csv.reader(io.StringIO(text)))
            text = "\n".join(" | ".join(cell.strip() for cell in row if cell.strip()) for row in rows)
        elif ext == "json":
            try:
                text = json.dumps(json.loads(text), indent=1, ensure_ascii=False)
            except ValueError:
                pass
        return ParsedDocument(pages=[Page(number=1, text=text)])


# ------------------------------------------------------------------ HTML


class _HtmlText(HTMLParser):
    BLOCK = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "header",
        "footer",
        "table",
        "pre",
        "blockquote",
    }
    SKIP = {"script", "style", "noscript", "template", "svg", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.SKIP:
            self._skip += 1
        if tag == "title":
            self._in_title = True
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        if tag == "title":
            self._in_title = False
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
            return
        if not self._skip:
            self.parts.append(data)


class HtmlReader(Reader):
    name = "html"
    extensions = ("html", "htm", "xhtml")

    def read(self, path: Path, settings: Settings) -> ParsedDocument:
        parser = _HtmlText()
        parser.feed(_read_text_file(path))
        return ParsedDocument(pages=[Page(number=1, text="".join(parser.parts))], title=parser.title.strip())


# ------------------------------------------------------------------ images


class ImageReader(Reader):
    """Text inside images, through OCR. Registered only when OCR is enabled."""

    name = "image-ocr"
    extensions = ("png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp")

    def read(self, path: Path, settings: Settings) -> ParsedDocument:
        from PIL import Image

        with Image.open(path) as image:
            image.load()
            text, confidence = ocr.image_to_text(image)
        return ParsedDocument(
            pages=[Page(number=1, text=text, ocr=True)], metadata={"ocr_confidence": round(confidence, 3)}
        )


def register_builtin_readers(settings: Settings) -> None:
    """Register the readers that the current settings allow."""
    register(PdfReader())
    register(DocxReader())
    register(PptxReader())
    register(TextReader())
    register(HtmlReader())
    if settings.ocr:
        if ocr.available():
            register(ImageReader())
        else:
            log.warning(
                "CSE_OCR is on but the ocr extra is missing; images are skipped",
                extra_fields={"reason": ocr.why_unavailable()},
            )
