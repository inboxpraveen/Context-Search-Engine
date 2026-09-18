"""Readers turn files into pages; failures become ReadError, never crashes."""

from __future__ import annotations

import pytest

from contextsearch.errors import ReadError, UnsupportedFileError
from contextsearch.readers import normalize_text, read_document, register_builtin_readers
from tests.conftest import make_pdf, make_settings


@pytest.fixture
def settings(tmp_path):
    s = make_settings(tmp_path)
    register_builtin_readers(s)
    return s


def test_pdf_pages_and_title(tmp_path, settings):
    path = tmp_path / "doc.pdf"
    path.write_bytes(make_pdf(["First page about refunds.", "Second page about shipping."], title="Refund guide"))
    doc = read_document(path, settings)
    assert [p.number for p in doc.pages] == [1, 2]
    assert "refunds" in doc.pages[0].text and "shipping" in doc.pages[1].text
    assert doc.title == "Refund guide"
    assert doc.reader == "pdf"


def test_docx_paragraphs_tables_and_heading_title(tmp_path, settings):
    from docx import Document

    path = tmp_path / "handbook.docx"
    d = Document()
    d.add_heading("Employee handbook", 1)
    d.add_paragraph("Leave requests need two weeks notice.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Leave", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "Annual", "25"
    d.add_paragraph("After the table.")
    d.save(str(path))
    doc = read_document(path, settings)
    text = doc.pages[0].text
    assert doc.title == "Employee handbook"
    assert text.index("Leave requests") < text.index("Annual | 25") < text.index("After the table")


def test_html_strips_scripts_and_keeps_title(tmp_path, settings):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>Team notes</title><style>b{}</style></head><body><h1>Hello</h1><p>Body text.</p><script>var x=1;</script></body></html>",
        encoding="utf-8",
    )
    doc = read_document(path, settings)
    assert doc.title == "Team notes"
    assert "Body text." in doc.pages[0].text
    assert "var x" not in doc.pages[0].text and "b{}" not in doc.pages[0].text


def test_markdown_title_drops_heading_marker(tmp_path, settings):
    path = tmp_path / "notes.md"
    path.write_text("# Refund policy\n\nBody.\n", encoding="utf-8")
    assert read_document(path, settings).title == "Refund policy"


def test_csv_rows_are_joined(tmp_path, settings):
    path = tmp_path / "prices.csv"
    path.write_text("sku,price\nA1,9.50\n", encoding="utf-8")
    assert "A1 | 9.50" in read_document(path, settings).pages[0].text


def test_unsupported_and_broken_files(tmp_path, settings):
    with pytest.raises(UnsupportedFileError):
        read_document(tmp_path / "x.exe", settings)
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a pdf at all")
    with pytest.raises(ReadError):
        read_document(broken, settings)


def test_normalize_text_collapses_whitespace_but_keeps_paragraphs():
    text = "a" + chr(0xA0) + " b  c\r\nd   \n\n\n\ne"  # includes a no-break space
    assert normalize_text(text) == "a b c\nd\n\ne"
